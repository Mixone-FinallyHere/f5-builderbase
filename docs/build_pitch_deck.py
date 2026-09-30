from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

C=lambda h: RGBColor.from_string(h)
BG,SURF,SURF2,BORDER,PRI,SEC,TXT,MUT="07080C","10121A","171A25","262A3A","7C5CFF","2DE2C4","F5F6FA","9AA0B4"
HF,BF="Space Grotesk","Inter"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]

def rect(s,x,y,w,h,fill,line=None,shape=MSO_SHAPE.RECTANGLE,radius=None):
    r=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb=C(fill)
    if line: r.line.color.rgb=C(line); r.line.width=Pt(1)
    else: r.line.fill.background()
    r.shadow.inherit=False
    if radius is not None and shape==MSO_SHAPE.ROUNDED_RECTANGLE: r.adjustments[0]=radius
    return r
def text(s,x,y,w,h,t,size=18,color=TXT,font=BF,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,spacing=None):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    lines=t if isinstance(t,list) else [t]
    for i,l in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        if spacing: p.space_after=Pt(spacing)
        r=p.add_run(); r.text=l; f=r.font; f.size=Pt(size); f.color.rgb=C(color); f.name=font; f.bold=bold
    return tb
def base(n,label,title):
    s=prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb=C(BG)
    rect(s,0,0,0.12,7.5,PRI)
    text(s,0.9,0.6,8,0.3,f"{n:02d}  ·  {label.upper()}",12,SEC,BF,True)
    text(s,0.9,1.0,11.5,1.2,title,40,TXT,HF,True)
    text(s,0.9,6.9,8,0.3,"F5 - Builderbase",11,MUT,BF)
    text(s,11.4,6.9,1.1,0.3,f"{n}/10",11,MUT,BF,align=PP_ALIGN.RIGHT)
    return s
def card(s,x,y,w,h,title,body,accent=SEC):
    rect(s,x,y,w,h,SURF,BORDER,MSO_SHAPE.ROUNDED_RECTANGLE,0.06)
    rect(s,x+0.3,y+0.3,0.6,0.08,accent,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.5)
    text(s,x+0.3,y+0.55,w-0.6,0.5,title,20,TXT,HF,True)
    text(s,x+0.3,y+1.1,w-0.6,h-1.3,body,15,MUT,BF)

# 1 Title
s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=C(BG)
g=rect(s,8.3,-1.5,7,7,PRI,shape=MSO_SHAPE.OVAL); g.fill.fore_color.rgb=C("1A1538")
g2=rect(s,10.3,3.8,5,5,SEC,shape=MSO_SHAPE.OVAL); g2.fill.fore_color.rgb=C("0C2A2B")
rect(s,0.9,1.6,1.4,0.42,SURF,BORDER,MSO_SHAPE.ROUNDED_RECTANGLE,0.5)
text(s,0.9,1.6,1.4,0.42,"HACKATHON 2026",10,MUT,BF,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
text(s,0.9,2.4,9,1.2,"F5 - Builderbase",72,TXT,HF,True)
text(s,0.9,3.7,8,0.5,"Pitch deck",36,PRI,HF,True)
text(s,0.9,4.5,8,0.8,"[Tagline: one sentence on who it's for and why it matters]",20,MUT,BF)
text(s,0.9,6.3,10,0.4,"[Team names]  ·  [Event name]  ·  [Date]",14,MUT,BF)
rect(s,0.9,5.65,1.2,0.06,SEC)

specs=[
 (2,"Hook","[One surprising stat, question, or 1-sentence story.]",None),
 (3,"Problem","[Who hurts, how often, what does it cost?]",[("Who","[Target user and their context]"),("Pain","[What goes wrong today, in concrete terms]"),("Cost","[Hours, dollars, or risk per week]")]),
 (4,"Solution","[X helps WHO do WHAT by HOW.]",[("Benefit 1","[Short, outcome-focused]"),("Benefit 2","[Short, outcome-focused]"),("Benefit 3","[Short, outcome-focused]")]),
 (5,"Demo","Show the golden path in 3 steps.",[("Step 1","[Screenshot / action]"),("Step 2","[Screenshot / action]"),("Step 3","[Result / wow moment]")]),
 (6,"How it works","[Architecture in one picture.]",[("Frontend","[e.g. Next.js, Tailwind]"),("Backend / AI","[APIs, models, data]"),("Built tonight","[What's new vs. reused]")]),
 (7,"Market","[Who pays, and how many of them?]",[("TAM","[$ / users]"),("SAM","[$ / users]"),("SOM","[$ / users, year 1]")]),
 (8,"Why now","[The shift that makes this possible today.]",[("Tech","[What just became possible]"),("Behavior","[What users now expect]"),("Timing","[Regulation, cost curve, trend]")]),
 (9,"Team","[Why this team wins.]",[("[Name]","[Role · one credential]"),("[Name]","[Role · one credential]"),("[Name]","[Role · one credential]")]),
]
for n,label,title,cards in specs:
    s=base(n,label,title)
    if n==2:
        rect(s,0.9,2.6,11.5,3.6,SURF,BORDER,MSO_SHAPE.ROUNDED_RECTANGLE,0.04)
        text(s,1.5,3.0,10.3,1.6,"[ 00% ]",96,PRI,HF,True)
        text(s,1.5,4.9,10.3,1.0,"[What the number means, in one line.]",22,MUT,BF)
    else:
        w=3.65
        for i,(t,b) in enumerate(cards):
            card(s,0.9+i*(w+0.28),2.7,w,3.4,t,b,[SEC,PRI,SEC][i])
# 10 Ask
s=base(10,"Ask","[What we want from you.]")
rect(s,0.9,2.7,7.2,3.4,SURF,BORDER,MSO_SHAPE.ROUNDED_RECTANGLE,0.05)
text(s,1.3,3.0,6.4,2.8,["[ ] Feedback / mentorship","[ ] Pilot users or design partners","[ ] Funding or credits: [amount]","[ ] Introductions to: [who]"],20,TXT,BF,spacing=12)
rect(s,8.4,2.7,4.0,3.4,PRI,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.05)
text(s,8.7,3.0,3.4,0.5,"Try it now",22,"FFFFFF",HF,True)
text(s,8.7,3.6,3.4,1.2,["[QR code]","[live URL]","github.com/WhiteChair/f5-builderbase"],14,"FFFFFF",BF,spacing=6)
prs.core_properties.title="F5 - Builderbase Pitch Template"
prs.core_properties.author="F5 - Builderbase"
prs.save("F5-Builderbase-Pitch-Template.pptx")
